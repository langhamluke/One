"""Import Raising Cane's weekly labor workbooks.

Each store keeps one Excel workbook per week (Wednesday through Tuesday).
What this module reads from it:

- LABOR TRACKER tab: "Projected Customers" and "Actual Customers" per day.
  The actual row is the only real demand signal in the workbook. A zero or
  blank actual means the week was not closed out; it is reported as missing,
  never filled with the projection.
- Day tabs (Weds, Thurs, Fri, Sat, Sun, Mon, Tues): the store's half-hour
  forecast of guests and sales, and the planned crew per station (Bird,
  Board, Toast, DT order, ...). Columns are located by their header labels,
  so older workbooks with shifted columns parse the same way. Rows after
  midnight are stored with a 1900-01-01 placeholder date; they belong to the
  same business day and are placed on the next calendar day.

The workbooks carry no dates. A week's dates come from the file name
("P2W1 2.5 - 2.11.xlsx", "WK3 P11 11.19-11.25.xlsx", ...) and the year from
the folder (data/2025/Q1/P2/...). Every store week starts on a Wednesday, so
a parsed start that is not a Wednesday is flagged. Files without a date
("P7W4.xlsx") take theirs from sibling weeks in the same period.

Nothing here sends data anywhere. The importer runs on the operator's own
machine and writes plain CSV files for the forecasting engine.
"""

from __future__ import annotations

import datetime as dt
import json
import re
import warnings
from collections import defaultdict
from dataclasses import dataclass, field
from pathlib import Path

import numpy as np
import pandas as pd

WEEK_DAYS = ["Wednesday", "Thursday", "Friday", "Saturday", "Sunday", "Monday", "Tuesday"]
DAY_SHEET_NAMES = {
    "Wednesday": ("Weds", "Wed", "Wednesday", "WED", "WEDS"),
    "Thursday": ("Thurs", "Thu", "Thur", "Thursday", "THURS"),
    "Friday": ("Fri", "Friday", "FRI"),
    "Saturday": ("Sat", "Saturday", "SAT"),
    "Sunday": ("Sun", "Sunday", "SUN"),
    "Monday": ("Mon", "Monday", "MON"),
    "Tuesday": ("Tues", "Tue", "Tuesday", "TUES"),
}
DAY_ALIASES = {
    "wednesday": "Wednesday", "weds": "Wednesday", "wed": "Wednesday",
    "thursday": "Thursday", "thurs": "Thursday", "thur": "Thursday", "thu": "Thursday",
    "friday": "Friday", "fri": "Friday",
    "saturday": "Saturday", "sat": "Saturday",
    "sunday": "Sunday", "sun": "Sunday",
    "monday": "Monday", "mon": "Monday",
    "tuesday": "Tuesday", "tues": "Tuesday", "tue": "Tuesday",
}
# Positional fallbacks seen in the wild, used only when header labels are missing.
FALLBACK_LAYOUTS = (
    {"time": 1, "sales": 2, "guests": 3, "crew": 20},
    {"time": 3, "sales": 4, "guests": 5, "crew": 24},
)
LATE_NIGHT_CUTOFF = dt.time(5, 0)
# LABOR TRACKER / WEEKLY RECAP rows read per day, by label prefix. The labor rows let
# the import put a number on what forecast misses cost: crew hours worked versus the
# hours the store's own labor matrix allows for the customers who actually came.
ROW_LABELS = {
    "projected customers": "projected",
    "actual customers": "actual",
    "projected labor target": "target_hours_projected",
    "projected allowable crew": "target_hours_projected",
    "scheduled crew hours": "scheduled_hours",
    "actual labor target": "target_hours_actual",
    "actual allowable crew": "target_hours_actual",
    "actual crew hours": "actual_hours",
}
DAILY_FIELDS = ("projected", "actual", "target_hours_projected", "scheduled_hours", "target_hours_actual", "actual_hours")


# --------------------------------------------------------------------------- file names

@dataclass
class WeekRef:
    path: Path
    year: int | None
    period: int | None
    week: int | None
    start: dt.date | None
    notes: list[str] = field(default_factory=list)

    @property
    def label(self) -> str:
        return self.path.name


_PERIOD = re.compile(r"(?i)P\s*(\d{1,2})(?!\d)")
_WEEK = re.compile(r"(?i)(?:^|[^A-Z])W(?:K|EEK)?\s*(\d)(?!\d)")
_DATE = re.compile(r"(\d{1,2})\.(\d{1,2})")
_YEAR = re.compile(r"^(20\d\d)$")


def parse_week_name(path: Path) -> WeekRef:
    """Pull year, period, week, and start date out of a workbook path."""
    stem = path.stem
    year = next((int(m.group(1)) for part in reversed(path.parent.parts) if (m := _YEAR.match(part))), None)
    period = int(m.group(1)) if (m := _PERIOD.search(stem)) else None
    week = int(m.group(1)) if (m := _WEEK.search(stem)) else None
    ref = WeekRef(path, year, period, week, None)
    m = _DATE.search(stem)
    if m and year:
        month, day = int(m.group(1)), int(m.group(2))
        y = year
        # Period 1 can start on Dec 31 of the prior year ("WK1 P1 12.31-1.6" filed under 2026).
        if month == 12 and (period == 1 or (period is None and week == 1)):
            y = year - 1
        try:
            ref.start = dt.date(y, month, day)
        except ValueError:
            ref.notes.append(f"unreadable start date {month}.{day}")
    elif not m:
        ref.notes.append("no date in file name")
    if ref.start and ref.start.weekday() != 2:
        # Snap to the nearest Wednesday; store weeks always start Wednesday.
        offset = (2 - ref.start.weekday() + 3) % 7 - 3
        snapped = ref.start + dt.timedelta(days=offset)
        ref.notes.append(f"start {ref.start} is not a Wednesday; using {snapped}")
        ref.start = snapped
    return ref


def resolve_missing_starts(refs: list[WeekRef]) -> None:
    """Infer start dates for files like 'P7W4.xlsx' from siblings in the same period."""
    known: dict[tuple[int, int], list[WeekRef]] = defaultdict(list)
    for r in refs:
        if r.start and r.year and r.period and r.week:
            known[(r.year, r.period)].append(r)
    for r in refs:
        if r.start or not (r.year and r.period and r.week):
            continue
        sib = known.get((r.year, r.period))
        if sib:
            s = min(sib, key=lambda x: abs(x.week - r.week))
            r.start = s.start + dt.timedelta(weeks=r.week - s.week)
            r.notes.append(f"start inferred as {r.start} from {s.label}")


def discover_workbooks(root: Path) -> tuple[list[WeekRef], list[tuple[Path, str]]]:
    """Find weekly workbooks under root. Returns (weeks, skipped-with-reason)."""
    refs, skipped = [], []
    for p in sorted(root.rglob("*.xlsx")):
        if p.name.startswith("~$"):
            skipped.append((p, "Excel lock file"))
            continue
        ref = parse_week_name(p)
        if ref.period is None and ref.week is None:
            skipped.append((p, "not a weekly workbook name"))
            continue
        refs.append(ref)
    resolve_missing_starts(refs)
    good, seen = [], {}
    for r in refs:
        if r.start is None:
            skipped.append((r.path, "could not determine week start: " + "; ".join(r.notes)))
        elif r.start in seen:
            skipped.append((r.path, f"duplicate of {seen[r.start].label} (week of {r.start})"))
        else:
            seen[r.start] = r
            good.append(r)
    return sorted(good, key=lambda r: r.start), skipped


# --------------------------------------------------------------------------- workbook parsing

def _norm(v) -> str:
    return re.sub(r"\s+", " ", str(v)).strip().lower() if v is not None else ""


def _num(v) -> float | None:
    if isinstance(v, bool) or v is None:
        return None
    if isinstance(v, (int, float)):
        return float(v) if np.isfinite(v) else None
    try:
        return float(str(v).replace(",", "").strip())
    except ValueError:
        return None


def _find_sheet(names: list[str], candidates: tuple[str, ...]) -> str | None:
    lower = {n.strip().lower(): n for n in names}
    for c in candidates:
        if c.lower() in lower:
            return lower[c.lower()]
    return None


def _daily_customers(rows: list[tuple]) -> dict[str, dict[str, float | None]]:
    """Projected and actual customers by weekday from a LABOR TRACKER / WEEKLY RECAP grid."""
    day_cols: dict[str, int] = {}
    for r in rows[:25]:
        found = {}
        for j, v in enumerate(r):
            name = DAY_ALIASES.get(_norm(v))
            if name and name not in found:
                found[name] = j
        if len(found) >= 6:
            day_cols = found
            break
    out: dict[str, dict[str, float | None]] = {d: dict.fromkeys(DAILY_FIELDS) for d in WEEK_DAYS}
    if not day_cols:
        return out
    done = set()
    for r in rows:
        label = next((_norm(v) for v in r[:4] if isinstance(v, str) and v.strip()), "")
        key = next((k for prefix, k in ROW_LABELS.items() if label.startswith(prefix)), None)
        if key and key not in done:
            for day, j in day_cols.items():
                out[day][key] = _num(r[j]) if j < len(r) else None
            done.add(key)
    return out


def _slot_columns(rows: list[tuple]) -> tuple[int, dict[str, int], dict[str, int]]:
    """Locate the header row and the time / sales / guests / crew / station columns."""
    for i, r in enumerate(rows[:8]):
        labels = {j: _norm(v) for j, v in enumerate(r) if isinstance(v, str) and v.strip()}
        time_c = next((j for j, s in labels.items() if s == "time"), None)
        guests_c = next((j for j, s in labels.items() if s == "guests"), None)
        if time_c is None or guests_c is None:
            continue
        sales_c = next((j for j, s in labels.items() if s.startswith("sales")), None)
        crew_c = next((j for j, s in labels.items() if "total crew" in s), None)
        stations = {}
        if crew_c is not None:
            for j in range(guests_c + 1, crew_c):
                if j in labels:
                    stations[str(r[j]).strip()] = j
        return i, {"time": time_c, "sales": sales_c, "guests": guests_c, "crew": crew_c}, stations
    # No header labels: pick the fallback layout whose time column holds the most times.
    best, best_n = FALLBACK_LAYOUTS[0], -1
    for lay in FALLBACK_LAYOUTS:
        n = sum(isinstance(r[lay["time"]], (dt.time, dt.datetime)) for r in rows if len(r) > lay["time"])
        if n > best_n:
            best, best_n = lay, n
    return 1, dict(best), {}


def _slot_timestamp(day: dt.date, v) -> pd.Timestamp | None:
    if isinstance(v, dt.datetime):
        # 1900-01-01 HH:MM is how Excel stores times past midnight in these sheets.
        if v.year <= 1900:
            return pd.Timestamp(dt.datetime.combine(day + dt.timedelta(days=1), v.time()))
        return pd.Timestamp(dt.datetime.combine(day, v.time()))
    if isinstance(v, dt.time):
        d = day + dt.timedelta(days=1) if v < LATE_NIGHT_CUTOFF else day
        return pd.Timestamp(dt.datetime.combine(d, v))
    return None


@dataclass
class ParsedWeek:
    ref: WeekRef
    daily: pd.DataFrame       # business_date, weekday, projected, actual
    slots: pd.DataFrame       # business_date, ts, guests, sales, crew
    deployment: pd.DataFrame  # business_date, ts, guests, station, crew
    notes: list[str]


def parse_workbook(ref: WeekRef) -> ParsedWeek:
    import openpyxl

    notes = list(ref.notes)
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        wb = openpyxl.load_workbook(ref.path, read_only=True, data_only=True)
    try:
        sheets = wb.sheetnames
        cust = {d: dict.fromkeys(DAILY_FIELDS) for d in WEEK_DAYS}
        for tab in ("LABOR TRACKER", "WEEKLY RECAP"):
            name = _find_sheet(sheets, (tab,))
            if not name:
                continue
            got = _daily_customers(list(wb[name].iter_rows(min_row=1, max_row=80, max_col=30, values_only=True)))
            for d in WEEK_DAYS:
                for k in DAILY_FIELDS:
                    if cust[d][k] is None and got[d][k] is not None:
                        cust[d][k] = got[d][k]
            if all(cust[d]["projected"] is not None and cust[d]["actual_hours"] is not None for d in WEEK_DAYS):
                break
        daily_rows, slot_rows, dep_rows = [], [], []
        for k, day in enumerate(WEEK_DAYS):
            bdate = ref.start + dt.timedelta(days=k)
            proj, act = cust[day]["projected"], cust[day]["actual"]
            sheet = _find_sheet(sheets, DAY_SHEET_NAMES[day])
            day_guests = 0.0
            if sheet:
                rows = list(wb[sheet].iter_rows(min_row=1, max_row=80, max_col=40, values_only=True))
                hdr, cols, stations = _slot_columns(rows)
                for r in rows[hdr + 1:]:
                    if len(r) <= cols["guests"]:
                        continue
                    ts = _slot_timestamp(bdate, r[cols["time"]])
                    if ts is None:
                        continue  # subtotal rows carry text in the time column
                    g = _num(r[cols["guests"]]) or 0.0
                    s = _num(r[cols["sales"]]) if cols.get("sales") is not None else None
                    c = _num(r[cols["crew"]]) if cols.get("crew") is not None and cols["crew"] < len(r) else None
                    slot_rows.append({"business_date": bdate, "ts": ts, "guests": g, "sales": s, "crew": c})
                    day_guests += g
                    for st, j in stations.items():
                        v = _num(r[j]) if j < len(r) else None
                        if v is not None:
                            dep_rows.append({"business_date": bdate, "ts": ts, "guests": g, "station": st, "crew": v})
            else:
                notes.append(f"no {day} tab")
            if proj and day_guests and abs(day_guests / proj - 1) > 0.05:
                notes.append(f"{day}: half-hour guests sum to {day_guests:.0f} vs projected {proj:.0f}")
            daily_rows.append({
                "business_date": bdate, "weekday": day,
                "projected": proj if proj and proj > 0 else None,
                "actual": act if act and act > 0 else None,
                **{k: (cust[day][k] if cust[day][k] and cust[day][k] > 0 else None)
                   for k in ("target_hours_projected", "scheduled_hours", "target_hours_actual", "actual_hours")},
                "forecast_sum": day_guests,
            })
    finally:
        wb.close()
    return ParsedWeek(ref, pd.DataFrame(daily_rows), pd.DataFrame(slot_rows), pd.DataFrame(dep_rows), notes)


# --------------------------------------------------------------------------- learned staffing

def learn_deployment_standards(dep: pd.DataFrame, band: int = 5, min_obs: int = 6) -> pd.DataFrame:
    """Crew per station as a step function of half-hour forecast guests.

    For each station and guest band (0-4, 5-9, ...), take the median planned
    crew across all weeks, then force it to be non-decreasing in guests so
    a busier half hour never gets fewer people. Bands with too few
    observations inherit from the band below.
    """
    if dep.empty:
        return pd.DataFrame(columns=["station", "guests_from", "guests_to", "crew", "observations"])
    d = dep[dep["guests"] > 0].copy()
    d["band"] = (d["guests"] // band).astype(int)
    out = []
    for st, g in d.groupby("station"):
        stats = g.groupby("band")["crew"].agg(["median", "count"])
        top = int(stats.index.max())
        crew_prev = 0.0
        for b in range(top + 1):
            if b in stats.index and stats.loc[b, "count"] >= min_obs:
                crew = max(float(stats.loc[b, "median"]), crew_prev)
                n = int(stats.loc[b, "count"])
            else:
                crew, n = crew_prev, int(stats.loc[b, "count"]) if b in stats.index else 0
            crew_prev = crew
            out.append({"station": st, "guests_from": b * band, "guests_to": (b + 1) * band - 1,
                        "crew": round(crew, 2), "observations": n})
    res = pd.DataFrame(out)
    # Drop stations the store never staffs.
    keep = res.groupby("station")["crew"].max()
    return res[res["station"].isin(keep[keep > 0].index)].reset_index(drop=True)


# --------------------------------------------------------------------------- import

@dataclass
class ImportResult:
    daily: pd.DataFrame
    slots: pd.DataFrame
    deployment: pd.DataFrame
    standards: pd.DataFrame
    report: dict


def import_workbooks(root: Path, open_hour: int = 10, close_hour: int = 24, progress=None) -> ImportResult:
    refs, skipped = discover_workbooks(root)
    if not refs:
        raise FileNotFoundError(f"No weekly workbooks found under {root}")
    dailies, slots, deps, week_notes, failed = [], [], [], {}, []
    for i, ref in enumerate(refs):
        if progress:
            progress(i + 1, len(refs), ref.label)
        try:
            with warnings.catch_warnings():
                warnings.simplefilter("ignore")  # openpyxl notes unsupported conditional formatting per file
                pw = parse_workbook(ref)
        except Exception as exc:  # noqa: BLE001 - one corrupt file must not sink the whole import
            failed.append((ref.path, f"{type(exc).__name__}: {exc}"))
            continue
        dailies.append(pw.daily.assign(file=ref.label, period=ref.period, week=ref.week))
        slots.append(pw.slots)
        deps.append(pw.deployment)
        if pw.notes:
            week_notes[ref.label] = pw.notes
    daily = pd.concat(dailies, ignore_index=True).sort_values("business_date")
    daily["business_date"] = pd.to_datetime(daily["business_date"])
    slot = pd.concat(slots, ignore_index=True) if slots else pd.DataFrame()
    dep = pd.concat(deps, ignore_index=True) if deps else pd.DataFrame()
    if not slot.empty:
        slot["business_date"] = pd.to_datetime(slot["business_date"])
    standards = learn_deployment_standards(dep)

    # Gaps: weeks absent between first and last week.
    starts = sorted({r.start for r in refs})
    expected = pd.date_range(starts[0], starts[-1], freq="7D").date
    missing_weeks = [d.isoformat() for d in expected if d not in set(starts)]
    closed = daily[(daily["projected"].isna()) & (daily["forecast_sum"] <= 0)]
    in_window = 0.0
    if not slot.empty:
        same_day = slot["ts"].dt.normalize() == slot["business_date"]
        hrs = slot["ts"].dt.hour
        in_window = float(slot.loc[same_day & (hrs >= open_hour) & (hrs < close_hour), "guests"].sum() / max(slot["guests"].sum(), 1))
    report = {
        "root": str(root),
        "weeks_imported": len(dailies),
        "first_week": starts[0].isoformat(),
        "last_week": starts[-1].isoformat(),
        "days": len(daily),
        "days_with_actuals": int(daily["actual"].notna().sum()),
        "days_missing_actuals": int((daily["actual"].isna() & daily["projected"].notna()).sum()),
        "days_closed": sorted(d.date().isoformat() for d in closed["business_date"]),
        "wednesdays_with_actuals": int(((daily["weekday"] == "Wednesday") & daily["actual"].notna()).sum()),
        "missing_weeks": missing_weeks,
        "share_of_guests_in_model_window": round(in_window, 4),
        "model_window": f"{open_hour:02d}:00-{close_hour:02d}:00 same calendar day",
        "stations_learned": sorted(standards["station"].unique().tolist()) if not standards.empty else [],
        "skipped_files": [(str(p.relative_to(root)) if p.is_relative_to(root) else str(p), why) for p, why in skipped],
        "failed_files": [(str(p), why) for p, why in failed],
        "week_notes": week_notes,
    }
    return ImportResult(daily, slot, dep, standards, report)


def write_data_dir(res: ImportResult, out: Path, open_hour: int = 10, close_hour: int = 24) -> dict:
    """Write the CSV layout the forecasting engine and web app read (see docs/05)."""
    out.mkdir(parents=True, exist_ok=True)
    daily = res.daily.copy()
    daily.rename(columns={"business_date": "date"}).to_csv(out / "daily_customers.csv", index=False, date_format="%Y-%m-%d")

    slot = res.slots.copy()
    same_day = slot["ts"].dt.normalize() == slot["business_date"]
    slot = slot[same_day & slot["ts"].dt.hour.between(open_hour, close_hour - 1)]
    slot["hour"] = slot["ts"].dt.floor("h")
    hourly = slot.groupby(["business_date", "hour"])[["guests", "sales"]].sum(min_count=1).reset_index()
    day_tot = slot.groupby("business_date")["guests"].sum()

    # Store forecast, hourly, for every imported day (including days without actuals).
    hourly[["hour", "guests"]].rename(columns={"hour": "ts", "guests": "store_forecast"}).to_csv(
        out / "store_forecast.csv", index=False, date_format="%Y-%m-%d %H:%M")

    # Hourly "transactions" for the engine: the day's actual customers spread over
    # hours in proportion to the store's own half-hour forecast. Daily totals are
    # real; the within-day shape is the store's forecast (no hourly actuals exist
    # in these workbooks). Days without actuals are left out entirely.
    act = daily.set_index("business_date")["actual"]
    full_day = res.slots.groupby("business_date")["guests"].sum()
    share_in_window = (day_tot / full_day.reindex(day_tot.index)).fillna(0)
    hourly["day_actual"] = hourly["business_date"].map(act)
    hourly["window_actual"] = hourly["day_actual"] * hourly["business_date"].map(share_in_window)
    hourly["share"] = hourly["guests"] / hourly["business_date"].map(day_tot)
    tx = hourly[hourly["day_actual"].notna() & hourly["share"].notna()].copy()
    tx["transactions"] = (tx["window_actual"] * tx["share"]).round(2)
    tx["net_sales"] = np.nan
    tx[["hour", "transactions", "net_sales"]].rename(columns={"hour": "ts"}).to_csv(
        out / "transactions.csv", index=False, date_format="%Y-%m-%d %H:%M")

    res.standards.to_csv(out / "labor_standards.csv", index=False)
    meta = {
        "source": "canes_weekly_workbooks",
        "hourly_shape": "store_forecast",
        "score_daily_only": True,
        "open_hour": open_hour,
        "close_hour": close_hour,
        "closed_dates": res.report["days_closed"],
        "imported_at": dt.datetime.now(dt.UTC).isoformat(timespec="seconds"),
    }
    (out / "meta.json").write_text(json.dumps(meta, indent=2))
    (out / "import_report.json").write_text(json.dumps(res.report, indent=2, default=str))
    return meta


def school_calendar_frame(start: dt.date, end: dt.date, ref_csv: Path | None = None) -> pd.DataFrame:
    """Expand the bundled St. Vrain calendar into one status per date (date,status)."""
    ref_csv = ref_csv or Path(__file__).resolve().parent.parent / "reference" / "st_vrain_school_calendar.csv"
    ranges = pd.read_csv(ref_csv, comment="#", parse_dates=["start", "end"])
    status = {}
    for r in ranges.itertuples(index=False):
        for d in pd.date_range(r.start, r.end, freq="D"):
            status[d.date()] = r.status
    rows = []
    for d in pd.date_range(start, end, freq="D"):
        dd = d.date()
        s = "weekend" if dd.weekday() >= 5 else status.get(dd, "in_session")
        rows.append({"date": dd.isoformat(), "status": s})
    return pd.DataFrame(rows)

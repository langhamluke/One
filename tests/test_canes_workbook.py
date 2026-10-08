"""Cane's weekly workbook import: file names, layouts, and the full app path.

Workbooks here are generated with the same layout as the real ones (tab
names, label rows, column headers, after-midnight encoding, subtotal rows)
but contain made-up numbers. No store data lives in this repository.
"""

import datetime as dt
import json
from pathlib import Path

import numpy as np
import openpyxl
import pandas as pd
import pytest

from flowcast import labor
from flowcast.connectors.canes_workbook import (
    WEEK_DAYS,
    discover_workbooks,
    import_workbooks,
    learn_deployment_standards,
    parse_week_name,
    parse_workbook,
    school_calendar_frame,
    write_data_dir,
)

DAY_TABS = ["Weds", "Thurs", "Fri", "Sat", "Sun", "Mon", "Tues"]
STATIONS = ["Bird", "Board", "DTOrd"]


def _after_midnight(h, m):
    """How these workbooks store times past midnight: Excel day 1 (1900-01-01), no time zone."""
    return dt.datetime(1900, 1, 1, h, m)  # noqa: DTZ001 - Excel cell values are naive by design


def _shape():
    """Half-hour share of a day's guests, 10:00-23:30 plus Fri/Sat late night."""
    slots = [dt.time(h, m) for h in range(10, 24) for m in (0, 30)]
    w = np.array([1, 2, 6, 9, 9, 7, 5, 4, 3, 3, 3, 4, 6, 8, 8, 7, 5, 4, 3, 3, 2, 2, 1, 1, 1, 1, 1, 1], float)
    return slots, w / w.sum()


def make_workbook(path: Path, projected, actual, layout="new", late_night=True, seed=0):
    rng = np.random.default_rng(seed)
    off = 0 if layout == "new" else 2
    wb = openpyxl.Workbook()
    lt = wb.active
    lt.title = "LABOR TRACKER"
    lt.cell(row=3, column=2, value="Canes 0000")
    for k, day in enumerate(WEEK_DAYS):
        lt.cell(row=4, column=4 + 2 * k, value=day)
        lt.cell(row=7, column=4 + 2 * k, value=projected[k])
        lt.cell(row=15, column=4 + 2 * k, value=(actual[k] if actual is not None else 0))
    lt.cell(row=7, column=2, value="Projected Customers")
    lt.cell(row=15, column=2, value="Actual Customers")
    # Labor rows: schedule built for the projection; target follows actual customers.
    for k in range(7):
        lt.cell(row=11, column=4 + 2 * k, value=round(projected[k] / 9.5, 1))
        if actual is not None:
            lt.cell(row=17, column=4 + 2 * k, value=round(actual[k] / 9.5, 1))
            lt.cell(row=18, column=4 + 2 * k, value=round(projected[k] / 9.5 * 0.6 + actual[k] / 9.5 * 0.4 + 2, 1))
    lt.cell(row=11, column=2, value="Scheduled Crew Hours")
    lt.cell(row=14, column=2, value="Actual Crew Deployment")
    lt.cell(row=17, column=2, value="Actual Labor Target")
    lt.cell(row=18, column=2, value="Actual Crew Hours")
    lt.cell(row=22, column=2, value="Actual vs Operator Forecast")
    slots, share = _shape()
    for k, tab in enumerate(DAY_TABS):
        ws = wb.create_sheet(tab)
        ws.cell(row=1, column=2 + off, value="Canes 0000 [ Note ]")
        for j, lab in enumerate(["Time", "Sales $", "Guests"] + STATIONS + ["Total Crew"]):
            ws.cell(row=2, column=2 + off + j, value=lab)
        r = 3
        guests = share * projected[k]
        night = tab in ("Fri", "Sat") and late_night
        if night:
            guests = guests * 0.97
        times = [dt.time(h, m) for h in range(5, 10) for m in (0, 30)] + slots
        vals = [0.0] * 10 + list(guests)
        if night:
            times += [_after_midnight(0, 0), _after_midnight(0, 30)]
            vals += [projected[k] * 0.02, projected[k] * 0.01]
        else:
            times += [_after_midnight(0, 0)]
            vals += [0.0]
        for t, g in zip(times, vals, strict=True):
            crew = [1 if g > 0 else 0, (1 + (g >= 40)) if g > 0 else 0, (1 + (g >= 60)) if g > 0 else 0]
            ws.cell(row=r, column=2 + off, value=t)
            ws.cell(row=r, column=3 + off, value=round(g * 14.5 + rng.normal(0, 1), 2))
            ws.cell(row=r, column=4 + off, value=round(g, 3))
            for j, c in enumerate(crew):
                ws.cell(row=r, column=5 + off + j, value=c)
            ws.cell(row=r, column=5 + off + len(STATIONS), value=sum(crew))
            r += 1
        ws.cell(row=r, column=2 + off, value="Total")
        ws.cell(row=r, column=4 + off, value=float(sum(vals)))
        ws.cell(row=r + 1, column=2 + off, value="Daily Total")
        ws.cell(row=r + 1, column=4 + off, value=float(sum(vals)) * 3)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


BASE = [1150, 1350, 1700, 1650, 1450, 1050, 1120]  # Wed..Tue


# --------------------------------------------------------------------------- file names

REAL_NAMES = {
    "2024/Q3/P7/P7W5 7.24 - 7.30.xlsx": "2024-07-24",
    "2024/Q4/P10/P10W1 9.25-10.1.xlsx": "2024-09-25",
    "2024/Q4/P12/P12W5 12.25 - 12.31.xlsx": "2024-12-25",
    "2025/Q1/P1/P1W1 1.1 - 1.7.xlsx": "2025-01-01",
    "2025/Q1/P1/P1W5 1.29 - 2.4.xlsx": "2025-01-29",
    "2025/Q3/P7/P7W1  7.2 - 7.8.xlsx": "2025-07-02",
    "2025/Q3/P8/P8 WK1 8.6-8.12.xlsx": "2025-08-06",
    "2025/Q4/P10/P10 WK 1 10.1-10.7.xlsx": "2025-10-01",
    "2025/Q4/P10/P10 WK5 10.29-11.04.xlsx": "2025-10-29",
    "2025/Q4/P11/WK 1 P11 11.5-11.11.xlsx": "2025-11-05",
    "2025/Q4/P11/Wk2 P11 11.12-11.8.xlsx": "2025-11-12",
    "2025/Q4/P11/WK3 P11 11.19-11.251.xlsx": "2025-11-19",
    "2025/Q4/P12/WK4 P12 12.24- 12.30.xlsx": "2025-12-24",
    "2026/Q1/P1/WK1 P1 12.31-1.6.xlsx": "2025-12-31",
    "2026/Q1/P2/WK4 P2 2.25-3.3.xlsx": "2026-02-25",
}


@pytest.mark.parametrize("rel,start", list(REAL_NAMES.items()))
def test_parse_real_file_name_patterns(rel, start):
    ref = parse_week_name(Path("data") / rel)
    assert ref.start == dt.date.fromisoformat(start)
    assert ref.start.weekday() == 2  # every store week starts on Wednesday
    assert ref.period is not None and ref.week is not None


def test_non_wednesday_start_is_snapped_and_flagged():
    ref = parse_week_name(Path("data/2025/Q1/P2/P2W1 2.6 - 2.12.xlsx"))
    assert ref.start == dt.date(2025, 2, 5) and any("not a Wednesday" in n for n in ref.notes)


def test_discovery_infers_undated_week_and_skips_strays(tmp_path):
    root = tmp_path / "data"
    make_workbook(root / "2024/Q3/P7/P7W5 7.24 - 7.30.xlsx", BASE, BASE)
    make_workbook(root / "2024/Q3/P7/P7W4.xlsx", BASE, BASE)
    make_workbook(root / "2024/Q4/P11/Daily Cleaning C0304.xlsx", BASE, BASE)
    (root / "2024/Q3/P7/~$P7W5.xlsx").write_bytes(b"lock")
    refs, skipped = discover_workbooks(root)
    starts = {r.path.name: r.start for r in refs}
    assert starts["P7W4.xlsx"] == dt.date(2024, 7, 17)
    assert {p.name for p, _ in skipped} == {"Daily Cleaning C0304.xlsx", "~$P7W5.xlsx"}


# --------------------------------------------------------------------------- workbook parsing

@pytest.mark.parametrize("layout", ["new", "old"])
def test_parse_workbook_both_layouts(tmp_path, layout):
    p = tmp_path / "data/2026/Q1/P2/WK1 P2 2.4-2.10.xlsx"
    actual = [b + 25 for b in BASE]
    make_workbook(p, BASE, actual, layout=layout)
    pw = parse_workbook(parse_week_name(p))
    d = pw.daily.set_index("weekday")
    assert list(pw.daily["business_date"]) == [dt.date(2026, 2, 4) + dt.timedelta(days=k) for k in range(7)]
    assert d.loc["Wednesday", "actual"] == actual[0]  # Wednesdays are read ("Weds" tab)
    assert d.loc["Saturday", "projected"] == BASE[3]
    assert d.loc["Saturday", "scheduled_hours"] == round(BASE[3] / 9.5, 1)
    assert d.loc["Saturday", "target_hours_actual"] == round(actual[3] / 9.5, 1)
    assert d.loc["Saturday", "actual_hours"] > 0
    # Subtotal rows are ignored: half-hour guests sum to the projection.
    assert np.allclose(pw.daily["forecast_sum"] / pw.daily["projected"], 1.0, atol=0.01)
    # After-midnight slots belong to the business day but sit on the next calendar day.
    fri = pw.slots[pw.slots["business_date"] == dt.date(2026, 2, 6)]
    late = fri[fri["ts"].dt.normalize() > pd.Timestamp("2026-02-06")]
    assert len(late) == 2 and late["guests"].sum() > 0
    assert set(pw.deployment["station"]) == set(STATIONS)


def test_missing_actuals_are_missing_not_forecast(tmp_path):
    p = tmp_path / "data/2026/Q1/P3/WK3 P3 3.18-3.24.xlsx"
    make_workbook(p, BASE, None)
    pw = parse_workbook(parse_week_name(p))
    assert pw.daily["actual"].isna().all()
    assert pw.daily["projected"].notna().all()


def test_learned_standards_are_monotone_and_usable(tmp_path):
    p = tmp_path / "data/2026/Q1/P2/WK1 P2 2.4-2.10.xlsx"
    make_workbook(p, BASE, BASE)
    std = learn_deployment_standards(parse_workbook(parse_week_name(p)).deployment, min_obs=1)
    for _, g in std.groupby("station"):
        assert g.sort_values("guests_from")["crew"].is_monotonic_increasing
    stds = labor.standards_from_frame(std)
    board = next(s for s in stds if s.station == "Board")
    assert board.need(20) <= board.need(140)  # 10 vs 70 guests per half hour
    plan = labor.staffing_plan(pd.DataFrame({"ts": pd.date_range("2026-02-04 10:00", periods=3, freq="h"),
                                             "forecast": [0.0, 40.0, 150.0]}), stds)
    assert plan.iloc[0]["total"] == 0 and plan.iloc[2]["total"] >= plan.iloc[1]["total"] > 0
    assert set(STATIONS) <= set(plan.columns)


def test_school_calendar_expansion():
    cal = school_calendar_frame(dt.date(2026, 3, 13), dt.date(2026, 3, 23)).set_index("date")["status"]
    assert cal["2026-03-16"] == "spring_break" and cal["2026-03-14"] == "weekend" and cal["2026-03-23"] == "in_session"


# --------------------------------------------------------------------------- import to app

def _make_history(root: Path, weeks: int = 12, missing_actual_week: int | None = None):
    start = dt.date(2025, 10, 1)
    rng = np.random.default_rng(1)
    for w in range(weeks):
        s = start + dt.timedelta(weeks=w)
        period = 10 + w // 4
        proj = [round(b * (1 + rng.normal(0, 0.04))) for b in BASE]
        act = None if w == missing_actual_week else [round(p * (1 + rng.normal(-0.02, 0.08))) for p in proj]
        name = f"P{period} WK{w % 4 + 1} {s.month}.{s.day}-{(s + dt.timedelta(days=6)).month}.{(s + dt.timedelta(days=6)).day}.xlsx"
        make_workbook(root / f"{s.year}/Q4/P{period}/{name}", proj, act, layout="old" if w < 4 else "new", seed=w)


def test_import_writes_data_dir(tmp_path):
    root = tmp_path / "data"
    _make_history(root, weeks=6, missing_actual_week=5)
    res = import_workbooks(root)
    out = tmp_path / "longmont"
    meta = write_data_dir(res, out)
    assert meta["score_daily_only"] is True and meta["hourly_shape"] == "store_forecast"
    rep = res.report
    assert rep["weeks_imported"] == 6 and rep["days_missing_actuals"] == 7
    assert rep["wednesdays_with_actuals"] == 5
    tx = pd.read_csv(out / "transactions.csv", parse_dates=["ts"])
    sf = pd.read_csv(out / "store_forecast.csv", parse_dates=["ts"])
    assert tx["ts"].dt.normalize().nunique() == 35  # days with actuals only
    assert sf["ts"].dt.normalize().nunique() == 42  # forecast for every day
    daily = pd.read_csv(out / "daily_customers.csv", parse_dates=["date"]).set_index("date")
    got = tx.groupby(tx["ts"].dt.normalize())["transactions"].sum()
    ratio = got / daily["actual"].reindex(got.index)
    assert ratio.between(0.95, 1.0001).all()  # late-night share sits outside the model window
    assert json.loads((out / "import_report.json").read_text())["stations_learned"] == sorted(STATIONS)


def test_imported_store_runs_in_the_app(tmp_path):
    from fastapi.testclient import TestClient

    from flowcast.app.config import Settings
    from flowcast.app.db import Database
    from flowcast.app.main import create_app
    from flowcast.app.state import StoreState
    from flowcast.synth import StoreConfig

    root = tmp_path / "data"
    _make_history(root, weeks=12, missing_actual_week=11)
    out = tmp_path / "longmont"
    res = import_workbooks(root)
    write_data_dir(res, out)
    school_calendar_frame(dt.date(2025, 9, 1), dt.date(2026, 2, 1)).to_csv(out / "school.csv", index=False)

    st = StoreState.build(StoreConfig(name="Test", state="CO"), data_dir=out, folds=4)
    # "Today" is the day after the last actual; the last week has only the store's forecast.
    assert st.today == pd.Timestamp("2025-12-17")
    assert st.week["store_forecast"].notna().all()
    assert st.stations == sorted(STATIONS) or set(st.stations) == set(STATIONS)
    assert st.data.score_daily_only
    assert "store_forecast" in st.selection.report.index
    # The store's forecast is the anchor: the chosen forecast is never worse than it out of sample.
    assert st.selection.anchor == "store_forecast"
    rep = st.selection.report
    assert rep.loc[st.selection.chosen, "daily_wape"] <= rep.loc["store_forecast", "daily_wape"] + 1e-12
    gap = st.labor_vs_target()
    assert gap is not None and gap["days"] >= 60
    assert set(gap["by_bucket"]["bucket"]) <= {"over-forecast >10%", "within 10%", "under-forecast >10%"}
    comp = st.daily_comparison()
    assert "Store's own forecast" in set(comp["forecast"])
    assert comp["days"].iloc[0] > 0

    settings = Settings(secret_key="y" * 40, db_path=tmp_path / "t.db", secure_cookies=False,
                        bootstrap_password="Admin-Pass-2026x", store_name="Test")
    app = create_app(settings, state=st, db=Database(settings.db_path))
    c = TestClient(app, follow_redirects=False)
    assert c.post("/login", data={"username": "admin", "password": "Admin-Pass-2026x"}).status_code == 303
    for path in ("/", "/forecast", "/people", "/ordering", "/assistant", "/admin"):
        r = c.get(path)
        assert r.status_code == 200, path
    home = c.get("/").text
    assert "store forecast" in home and "How each forecast did" in home
    assert "What forecast misses cost in labor" in home
    assert "DTOrd" in c.get("/people").text


def test_labor_vs_target_attributes_hours_to_forecast_misses():
    idx = pd.date_range("2026-01-01", periods=60, freq="D")
    rng = np.random.default_rng(0)
    actual = rng.integers(1000, 1600, len(idx)).astype(float)
    projected = actual * np.where(np.arange(len(idx)) % 5 == 0, 1.2, 1.0)  # every 5th day over-forecast 20%
    target = actual / 9.5
    scheduled = projected / 9.5
    worked = scheduled - 0.3 * (scheduled - target)  # managers claw back 30% in-day
    df = pd.DataFrame({"projected": projected, "actual": actual, "scheduled_hours": scheduled,
                       "actual_hours": worked, "target_hours_actual": target}, index=idx)
    g = labor.labor_vs_target(df, wage=16.0)
    bb = g["by_bucket"].set_index("bucket")
    assert abs(bb.loc["within 10%", "remaining_gap"]) < 1e-9
    over = bb.loc["over-forecast >10%"]
    assert over["planning_gap"] > 0 and over["adjustment"] < 0 and over["remaining_gap"] > 0
    assert g["excess_hours_on_over_forecast_days"] > 0 and g["excess_dollars_per_year"] > 0
    assert labor.labor_vs_target(df.drop(columns="actual_hours")) is None
